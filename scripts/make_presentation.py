#!/usr/bin/env python3
"""Генерирует docs/presentation.pptx — питч «Арена Переговоров» (ТЗ хакатона, п.4).

Фирменные цвета взяты из dist/css/app.css (тёмная тема). На каждом слайде —
заметки докладчика (notes). PDF-копия: `docs/presentation.pdf`.

Запуск:
    pip install python-pptx
    python3 scripts/make_presentation.py
    # при желании обновить PDF:
    soffice --headless --convert-to pdf --outdir docs docs/presentation.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ── Палитра (dist/css/app.css, dark theme) ─────────────────────────
BG = RGBColor(0x0A, 0x0E, 0x13)
SURFACE = RGBColor(0x13, 0x1A, 0x24)
SURFACE_2 = RGBColor(0x1A, 0x23, 0x31)
TEXT = RGBColor(0xE8, 0xEE, 0xF7)
MUTED = RGBColor(0x9A, 0xA7, 0xB8)
ACCENT = RGBColor(0xC6, 0xF2, 0x4E)
DANGER = RGBColor(0xF2, 0x6C, 0x6C)

FONT = "Segoe UI"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


# ── Базовые помощники ──────────────────────────────────────────────
def set_bg(slide, color=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def style_run(run, size, color=TEXT, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def add_para(tf, marker, text, size, *, color=TEXT, marker_color=ACCENT,
             bold=False, space_after=8, align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    if marker:
        r1 = p.add_run()
        r1.text = f"{marker} "
        style_run(r1, size, marker_color, bold=True)
    r2 = p.add_run()
    r2.text = text
    style_run(r2, size, color, bold=bold)
    return p


def new_slide(title=None, kicker=None, page=None, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_bg(slide)

    if title:
        if kicker:
            _, tf = add_textbox(slide, Inches(0.75), Inches(0.38), Inches(11.8), Inches(0.35))
            add_para(tf, None, kicker.upper(), 12, color=ACCENT, bold=True,
                     space_after=0, first=True)
            title_top = Inches(0.72)
        else:
            title_top = Inches(0.5)
        _, tf = add_textbox(slide, Inches(0.75), title_top, Inches(11.8), Inches(0.8))
        # Кегль по длине: длинный заголовок не должен переноситься на вторую
        # строку и наезжать на акцентную черту под ним (запас на подмену шрифта
        # в рендерах без Segoe UI).
        title_size = 30
        if len(title) > 42:
            title_size = 26
        if len(title) > 56:
            title_size = 22
        add_para(tf, None, title, title_size, color=TEXT, bold=True,
                 space_after=0, first=True)
        # Акцентная черта под заголовком
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.78), title_top + Inches(0.72),
            Inches(1.4), Inches(0.05),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = ACCENT
        bar.line.fill.background()

    # Футер
    _, tf = add_textbox(slide, Inches(0.75), Inches(7.02), Inches(9.0), Inches(0.35))
    add_para(tf, None, "Арена Переговоров · хакатон 2026", 10,
             color=MUTED, marker_color=MUTED, space_after=0, first=True)
    if page is not None:
        _, tf = add_textbox(slide, Inches(12.3), Inches(7.02), Inches(0.7), Inches(0.35))
        add_para(tf, None, str(page), 10, color=MUTED, marker_color=MUTED,
                 space_after=0, align=PP_ALIGN.RIGHT, first=True)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_card(slide, left, top, width, height, title, items, *,
             title_color=ACCENT, fill=SURFACE, item_size=15, marker="•"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = SURFACE_2
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.28)
    tf.margin_right = Inches(0.28)
    tf.margin_top = Inches(0.22)
    tf.margin_bottom = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    add_para(tf, None, title, 17, color=title_color, bold=True,
             space_after=10, first=True)
    for item in items:
        if isinstance(item, tuple):
            text, sub = item
        else:
            text, sub = item, False
        add_para(tf, None if sub else marker, text,
                 item_size - 2 if sub else item_size,
                 color=MUTED if sub else TEXT,
                 marker_color=MUTED if sub else ACCENT,
                 space_after=6)
    return shape


# ── 1. Титул ───────────────────────────────────────────────────────
s = new_slide(page=None, notes=(
    "Питч: «Арена Переговоров» — интерактивный симулятор переговоров. "
    "Демо: настройка контекста администратором → диалог → отчёт с обратной связью."))
_, tf = add_textbox(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(1.4))
add_para(tf, None, "«Арена Переговоров»", 54, color=TEXT, bold=True,
         space_after=0, first=True)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(3.25),
                         Inches(2.2), Inches(0.07))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
_, tf = add_textbox(s, Inches(0.9), Inches(3.55), Inches(11.5), Inches(1.6))
add_para(tf, None,
         "Интерактивный симулятор переговоров с ИИ-собеседником",
         24, color=ACCENT, bold=True, space_after=14, first=True)
add_para(tf, None,
         "Безопасная практика делового общения · исход зависит от вашей тактики · "
         "обратная связь после каждого диалога",
         17, color=MUTED, space_after=0)
_, tf = add_textbox(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.5))
add_para(tf, None, "MVP готов · Rust + SQLite + SPA · демо работает без интернета",
         13, color=MUTED, space_after=0, first=True)

# ── 2. Проблема и аудитория ────────────────────────────────────────
s = new_slide("Проблема и аудитория", "Зачем это нужно", page=1, notes=(
    "Очные тренинги работают, но не масштабируются: дорого, мало попыток, "
    "нет разбора именно ваших формулировок. Чтение теории без практики не работает."))
add_card(s, Inches(0.75), Inches(1.7), Inches(5.95), Inches(4.9),
         "Проблема", [
             "Переговоры учат «на бумаге»: лекции, чек-листы, кейсы без живого диалога",
             "Очные ролевые игры плохо масштабируются: дорого, не в любой момент, мало попыток",
             "Разбирают кейс «в общем», а не ваши конкретные формулировки",
             "Обратной связи после диалога почти нет: непонятно, что сработало",
         ], item_size=16)
add_card(s, Inches(7.05), Inches(1.7), Inches(5.55), Inches(4.9),
         "Целевая аудитория", [
             "Основная: студенты, стажёры, junior-специалисты — подготовка к реальным сделкам",
             "Вторичная: sales, закупки, account management — сложные кейсы и ультиматумы",
             "Канал: преподаватели и корпоративные тренеры — повторяемый симулятор с разбором",
             ("Не для: юридически значимых консультаций и замены коуча 1:1", True),
         ], item_size=16)

# ── 3. Ценность ────────────────────────────────────────────────────
s = new_slide("Чем это лучше обычного обучения", "Ценность", page=2, notes=(
    "Ключевой слайд продуктовой части: безопасная практика, многократные попытки, "
    "персональный разбор, настройка под задачу."))
_, tf = add_textbox(s, Inches(0.75), Inches(1.75), Inches(11.9), Inches(4.9))
rows = [
    ("Теория без тренировки", "живой диалог с ИИ-собеседником и неограниченные попытки"),
    ("Разбор кейса «на бумаге»", "скоринг каждой реплики: стратегия, аргументация, тон, SPIN"),
    ("Занятие по расписанию тренера", "сценарий под нужную задачу настраивается за минуты"),
    ("«Что я сделал не так?» — непонятно", "итоговый отчёт: финал, разбор и конкретные рекомендации"),
    ("Один кейс на всех", "один тренажёр под sales, закупки, HR и учебные программы"),
]
first = True
for left_text, right_text in rows:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(16)
    r = p.add_run(); r.text = left_text
    style_run(r, 17, MUTED, italic=True)
    r = p.add_run(); r.text = "   →   "
    style_run(r, 17, ACCENT, bold=True)
    r = p.add_run(); r.text = right_text
    style_run(r, 17, TEXT, bold=True)

# ── 4. Путь пользователя ───────────────────────────────────────────
s = new_slide("Путь пользователя: от входа до разбора", "UX", page=3, notes=(
    "Полный цикл проходится за одну сессию (десятки минут). "
    "Шаг 4 — ветвление: можно откатиться и попробовать другую тактику."))
steps = [
    ("1", "Вход", "регистрация / логин, язык ru·en"),
    ("2", "Каталог сценариев", "фильтры: сложность, сфера, поиск"),
    ("3", "Диалог", "текстом или голосовым звонком (STT/TTS)"),
    ("4", "Ветвление", "откат к реплике и форк — сравни тактики"),
    ("5", "Отчёт", "балл, финал, разбор, рекомендации"),
    ("6", "Прогрессия", "XP, уровень, лидерборд"),
    ("7", "Повтор", "новая попытка или другой сценарий"),
]
x = 0.75
for num, title, desc in steps:
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(x), Inches(2.1), Inches(1.66), Inches(3.4))
    card.fill.solid()
    card.fill.fore_color.rgb = SURFACE
    card.line.color.rgb = SURFACE_2
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.2)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    add_para(tf, None, num, 30, color=ACCENT, bold=True, space_after=6,
             align=PP_ALIGN.CENTER, first=True)
    add_para(tf, None, title, 15, color=TEXT, bold=True, space_after=8,
             align=PP_ALIGN.CENTER)
    add_para(tf, None, desc, 12, color=MUTED, space_after=0,
             align=PP_ALIGN.CENTER)
    x += 1.78
_, tf = add_textbox(s, Inches(0.75), Inches(5.8), Inches(11.9), Inches(0.8))
add_para(tf, None,
         "Одна сессия = полный цикл «вход → разбор». Попытки не ограничены — "
         "тренировка дешевле, чем ошибка в реальных переговорах.",
         15, color=MUTED, space_after=0, first=True)

# ── 5. Механики ────────────────────────────────────────────────────
s = new_slide("Механики: как формулировки превращаются в исход", "Игровой процесс",
              page=4, notes=(
    "Главное для жюри: исход действительно зависит от действий игрока — "
    "это детерминированный скоринг, а не «LLM решила, что вы проиграли»."))
add_card(s, Inches(0.75), Inches(1.7), Inches(5.95), Inches(4.95),
         "От стратегии к баллу", [
             "Каждая реплика анализируется: сотрудничество / компромисс / конфронтация, аргументация, тон",
             "Бонусы за техники: SPIN, интересы сторон, объективные критерии (Гарвардский метод)",
             "LLM-судья 0–10 дополняет эвристику (доля настраивается админом)",
             "Сложность easy / medium / hard меняет поведение собеседника",
             "BATNA обеих сторон заложен в сценарий: на hard собеседник использует свою альтернативу как рычаг",
         ], item_size=15)
add_card(s, Inches(7.05), Inches(1.7), Inches(5.55), Inches(4.95),
         "Ветвление и финалы", [
             "Дерево реплик: откат к реплике собеседника → форк (до 16 веток), прежний путь сохраняется",
             "Снапшот счёта в точке ветвления — можно сравнивать тактики",
             "Финалы по порогам балла: победа / компромисс / срыв переговоров",
             "Финалы описывает администратор — свой исход под каждую задачу",
         ], item_size=15)

# ── 6. Конфигурируемость ──────────────────────────────────────────
s = new_slide("Конфигурируемость под контекст",
              "Администратор", page=5, notes=(
    "Демонстрационный слайд: Админка → Сценарии → ИИ-генератор: "
    "задаём сферу, тему, сложность, роли, цели и тон → получаем черновик."))
add_card(s, Inches(0.75), Inches(1.7), Inches(5.95), Inches(4.95),
         "Контекст на входе", [
             "Сфера и тема переговоров",
             "Сложность: easy / medium / hard",
             "Роль и цели сторон, BATNA",
             "Тон, манера и черты собеседника",
             ("Это же влияет на system prompt диалога — сценарий и поведение ИИ связаны", True),
         ], item_size=16)
add_card(s, Inches(7.05), Inches(1.7), Inches(5.55), Inches(4.95),
         "Два пути получить сценарий", [
             "Конструктор — все поля вручную, свои финалы и пороги",
             "ИИ-генератор — бриф + контекст → готовый черновик за секунды",
             "Заданные поля применяются буквально — воля администратора приоритетнее ИИ",
             "Черновик вычитывается и публикуется одной кнопкой",
             "Импорт / экспорт JSON — обмен сценариями между группами",
         ], item_size=16)

# ── 7. Геймификация и обратная связь ──────────────────────────────
s = new_slide("Геймификация и обратная связь", "Опционально — но уже есть",
              page=6, notes=(
    "Очки и прогрессия поощряют повторять; отчёт превращает попытку в урок."))
add_card(s, Inches(0.75), Inches(1.7), Inches(3.87), Inches(4.9),
         "Очки и прогресс", [
             "Балл за каждый ход — видно сразу",
             "XP за завершённые сессии",
             "10 уровней с названиями",
             "Ледерборд между игроками",
         ], item_size=15)
add_card(s, Inches(4.83), Inches(1.7), Inches(3.87), Inches(4.9),
         "Циклы обратной связи", [
             "Реакция собеседника в процессе",
             "Итоговый отчёт: финал и баллы по измерениям",
             "Разбор (markdown) на языке интерфейса",
             "Персональные рекомендации: что улучшить",
         ], item_size=15)
add_card(s, Inches(8.89), Inches(1.7), Inches(3.71), Inches(4.9),
         "Живость диалога", [
             "Текстовый чат или голосовой звонок",
             "Реплика открытия из сценария",
             "Тон собеседника зависит от контекста",
             "Многократные попытки без ограничений",
         ], item_size=15)

# ── 8. Технические решения ────────────────────────────────────────
s = new_slide("Технические решения", "MVP", page=7, notes=(
    "Ключевое: прототип запускается на обычной машине за минуту; "
    "демо работает без интернета через встроенный mock-провайдер."))
add_card(s, Inches(0.75), Inches(1.7), Inches(5.95), Inches(4.95),
         "Архитектура", [
             "Rust / Axum — модульный монолит: domain · application · infrastructure · web",
             "SQLite с версионными миграциями — демо не ломается после перезапуска",
             "SPA vanilla-JS без сборщика — отдаётся сервером как есть",
             "HTTP API версионирован: /api/v1, health /health",
             "Docker-образ и docker-compose с healthcheck",
         ], item_size=15)
add_card(s, Inches(7.05), Inches(1.7), Inches(5.55), Inches(4.95),
         "AI, голос и безопасность", [
             "LLM не обязателен: OpenAI-совместимые / Anthropic / Gemini / Ollama",
             "Встроенный mock-провайдер — полное демо без ключей и интернета",
             "Голос опционально: STT и TTS (облачные и локальные модели)",
             "JWT + refresh, RBAC, Argon2, AES-256-GCM для API-ключей, rate-limit, аудит",
             "258 автотестов, CI: fmt + clippy -D warnings + test",
         ], item_size=15)

# ── 9. Границы MVP ────────────────────────────────────────────────
s = new_slide("Границы MVP и дорожная карта", "Честный scope", page=8, notes=(
    "Показываем, что уже работает, а что — осознанно вне MVP."))
add_card(s, Inches(0.75), Inches(1.7), Inches(5.95), Inches(4.95),
         "Есть в MVP", [
             "Диалог (текст и голос), скоринг ходов + LLM-судья",
             "Ветвление диалога, финалы по порогам, отчёт с рекомендациями",
             "Админка: сценарии, пользователи, провайдеры, настройки, аудит",
             "XP / уровни / лидерборд, i18n ru·en, mock-демо, Docker, CI",
         ], item_size=15)
add_card(s, Inches(7.05), Inches(1.7), Inches(5.55), Inches(4.95),
         "Дальше", [
             "Публичный деплой (Vercel / Railway / Render)",
             "Достижения и стрики — глубже геймификация",
             "Сюжетные деревья сценариев (ветки внутри кейса)",
             "Стриминг ответов, мобильная версия, командные режимы",
         ], title_color=MUTED, item_size=15)

# ── 10. Демо ──────────────────────────────────────────────────────
s = new_slide("Демонстрация: запуск за 3 шага", "Как показать жюри",
              page=9, notes=(
    "Демо-скрипт: 1) запуск, 2) вход admin/admin123, "
    "3) ИИ-генератор с контекстом → диалог → отчёт."))
_, tf = add_textbox(s, Inches(0.75), Inches(1.75), Inches(11.9), Inches(3.4))
demo_rows = [
    ("1.", "git clone <репозиторий> && cp .env.example .env && cargo run",
     "или docker compose up -d — без ручной настройки"),
    ("2.", "Открыть http://localhost:3001 → войти admin / admin123",
     "демо работает офлайн: встроенный mock-LLM отвечает без ключей"),
    ("3.", "Админка → Сценарии → ✨ ИИ-генератор (сфера, тема, сложность, роли, тон) → Игрок → Начать → Завершить → Отчёт",
     "настройка контекста → диалог → обратная связь — ключевой сценарий демо"),
]
first = True
for num, main, sub in demo_rows:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(6)
    r = p.add_run(); r.text = f"{num} "
    style_run(r, 19, ACCENT, bold=True)
    r = p.add_run(); r.text = main
    style_run(r, 17, TEXT, bold=True)
    p2 = tf.add_paragraph()
    p2.space_after = Pt(18)
    p2.level = 1
    r = p2.add_run(); r.text = f"     {sub}"
    style_run(r, 14, MUTED)
add_card(s, Inches(0.75), Inches(5.3), Inches(11.85), Inches(1.3),
         "Требования к машине жюри", [
             "Любой ноутбук с браузером · Rust или Docker · интернет не обязателен для демо",
         ], item_size=14)

# ── 11. Итог ──────────────────────────────────────────────────────
s = new_slide(None, page=None, notes=(
    "Финал: практика вместо теории, объяснимый результат, настройка под аудиторию. "
    "Ссылки: репозиторий, документация docs/, презентация."))
_, tf = add_textbox(s, Inches(0.9), Inches(1.4), Inches(11.5), Inches(1.0))
add_para(tf, None, "Итог", 40, color=TEXT, bold=True, space_after=0, first=True)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(2.35),
                         Inches(1.8), Inches(0.06))
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background()
_, tf = add_textbox(s, Inches(0.9), Inches(2.75), Inches(11.5), Inches(2.6))
for text in [
    "Практика вместо теории: безопасный диалог и неограниченные попытки",
    "Исход зависит от вашей тактики — и всегда объясняется разбором",
    "Настраивается под любую аудиторию: sales, закупки, HR, обучение",
]:
    add_para(tf, "▸", text, 19, color=TEXT, space_after=14,
             first=(text.startswith("Практика")))
_, tf = add_textbox(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(1.0))
add_para(tf, None,
         "Репозиторий · docs/CONCEPT.md · docs/DOCUMENTATION.md · docs/presentation.pptx",
         14, color=MUTED, space_after=6, first=True)
add_para(tf, None, "Спасибо! Вопросы?", 18, color=ACCENT, bold=True, space_after=0)

OUT = str(Path(__file__).resolve().parent.parent / "docs" / "presentation.pptx")
prs.save(OUT)
print(f"saved: {OUT} ({len(prs.slides)} slides)")
